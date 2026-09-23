import asyncio
import logging
import os
import base64
import json
import re
import io
import tempfile
from io import BytesIO
from datetime import datetime
import asyncpg
import edge_tts
import aiohttp
import requests
from huggingface_hub import InferenceClient
from pexafy import Client as PexafyClient
from aiogram import Bot, Dispatcher, types, BaseMiddleware, F
from aiogram.filters import Command
from aiogram.types import FSInputFile, BufferedInputFile, InlineKeyboardMarkup, InlineKeyboardButton
from openai import AsyncOpenAI
from aiohttp import web
from PIL import Image

# === PPTX ИМПОРТЫ — ГЛОБАЛЬНО ===
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# --- Конфиг ---
BOT_TOKEN = os.getenv("BOT_TOKEN")
GROQ_KEY = os.getenv("GROQ_API_KEY")
DATABASE_URL = os.getenv("DATABASE_URL")
HF_TOKEN = os.getenv("HF_TOKEN")
PEXAFY_KEY = os.getenv("PEXAFY_API_KEY")

client = AsyncOpenAI(base_url="https://api.groq.com/openai/v1", api_key=GROQ_KEY)
hf_client = InferenceClient(token=HF_TOKEN) if HF_TOKEN else None
pexafy_client = PexafyClient(PEXAFY_KEY) if PEXAFY_KEY else None

bot = Bot(token=BOT_TOKEN)
dp = Dispatcher()

ALLOWED_IDS = [5264513480, 8834374199, 5389046699, 2083728480, 6612130539]

db_pool = None
HISTORY_LIMIT = 20
MAX_CONTEXT_CHARS = 15000
CODE_EXTENSIONS = ["html", "py", "js", "css", "java", "cpp", "sql", "json"]
MAX_AUTO_PARTS = 10

# ============================================================
# ПАЛИТРЫ
# ============================================================
PALETTES = {
    "warm":   {"bg": (255, 248, 240), "accent": (224, 122, 95),  "text": (61, 64, 91),   "light": (242, 204, 143)},
    "cool":   {"bg": (240, 244, 248), "accent": (61, 90, 128),   "text": (41, 50, 65),   "light": (152, 193, 217)},
    "nature": {"bg": (244, 249, 244), "accent": (45, 106, 79),   "text": (27, 67, 50),   "light": (149, 213, 178)},
    "dark":   {"bg": (26, 26, 46),    "accent": (233, 69, 96),   "text": (255, 255, 255),"light": (22, 33, 62)},
    "purple": {"bg": (245, 240, 255), "accent": (124, 58, 237),  "text": (45, 45, 58),   "light": (196, 181, 253)},
    "pink":   {"bg": (255, 240, 245), "accent": (219, 39, 119),  "text": (45, 45, 58),   "light": (249, 168, 212)},
    "gold":   {"bg": (255, 251, 235), "accent": (217, 119, 6),   "text": (41, 37, 36),   "light": (252, 211, 77)},
}

def pick_palette(topic: str) -> dict:
    low = topic.lower()
    if any(w in low for w in ["космос", "звезд", "галактик", "планет", "астроном", "темн"]):
        return PALETTES["dark"]
    if any(w in low for w in ["кошк", "собак", "животн", "природа", "цвет", "сад", "лес"]):
        return PALETTES["warm"]
    if any(w in low for w in ["мор", "океан", "вод", "неб", "холод", "зим"]):
        return PALETTES["cool"]
    if any(w in low for w in ["эколог", "растен", "биолог", "здоров", "медицин"]):
        return PALETTES["nature"]
    if any(w in low for w in ["искусств", "музык", "поэз", "любов", "роман", "девуш", "цвет"]):
        return PALETTES["pink"]
    if any(w in low for w in ["истор", "деньг", "бизнес", "золот", "богат"]):
        return PALETTES["gold"]
    return PALETTES["purple"]

# ============================================================
# УТИЛИТЫ
# ============================================================
def split_message(text: str, limit: int = 4000) -> list[str]:
    if len(text) <= limit: return [text]
    parts = []
    while text:
        if len(text) <= limit:
            parts.append(text); break
        chunk = text[:limit]
        sp = chunk.rfind("\n")
        if sp == -1: sp = chunk.rfind(" ")
        if sp == -1: sp = limit
        parts.append(text[:sp])
        text = text[sp:].lstrip("\n")
    return parts

def trim_history_by_chars(history: list[dict], max_chars: int = MAX_CONTEXT_CHARS) -> list[dict]:
    total = 0; trimmed = []
    for msg in reversed(history):
        c = msg["content"]
        l = len(c) if isinstance(c, str) else sum(len(p.get("text", "")) + len(p.get("image_url", {}).get("url", "")) for p in c)
        if total + l > max_chars: break
        trimmed.append(msg); total += l
    return list(reversed(trimmed))

def detect_extension(text: str) -> str:
    if "<!DOCTYPE html>" in text or "<html" in text.lower(): return "html"
    if "def " in text and "import " in text: return "py"
    if "function " in text or "const " in text or "let " in text: return "js"
    if "#include" in text: return "cpp"
    if "SELECT " in text and "FROM " in text: return "sql"
    if text.strip().startswith("{") and text.strip().endswith("}"): return "json"
    return "txt"

def clean_code(text: str) -> str:
    text = text.replace("[file content end]", "").replace("[file content begin]", "")
    text = text.replace("// (продолжение следует)", "").replace("// (код готов)", "")
    text = text.replace("(продолжение следует)", "").replace("(код готов)", "")
    text = text.replace("```html", "").replace("```javascript", "").replace("```js", "")
    text = text.replace("```python", "").replace("```py", "").replace("```css", "")
    text = text.replace("```", "")
    return text.strip()

def extract_code(answer: str) -> str:
    if "```" in answer:
        matches = re.findall(r'```(?:\w+)?\n(.*?)```', answer, re.DOTALL)
        if matches: return clean_code("\n\n".join(matches))
    markers = ["<!DOCTYPE", "<html", "def ", "import ", "function ", "const ", "class ", "public class"]
    lines = answer.split("\n")
    start_idx = -1
    for i, line in enumerate(lines):
        if any(m in line for m in markers):
            start_idx = i; break
    if start_idx >= 0: return clean_code("\n".join(lines[start_idx:]))
    return clean_code(answer)

def is_code_complete(answer: str) -> bool:
    low = answer.lower()
    if "код готов" in low or "// (готово)" in low: return True
    if "продолжение следует" in low or "to be continued" in low: return False
    if "</html>" in low and "</script>" in low: return True
    if answer.count("```") >= 2 and "</html>" in answer: return True
    if answer.count("{") == answer.count("}") and answer.count("(") == answer.count(")") and len(answer) > 2000:
        return True
    return False

def parse_json_safe(raw: str):
    if not raw: return None
    result = None
    if "```" in raw:
        for p in raw.split("```"):
            p = p.strip()
            if p.startswith("json"): p = p[4:].strip()
            if p.startswith("[") or p.startswith("{"):
                try: result = json.loads(p); break
                except: continue
    if result is None:
        try: result = json.loads(raw.strip())
        except:
            m = re.search(r'\[.*\]', raw, re.DOTALL)
            if m:
                try: result = json.loads(m.group(0))
                except: return None
    if isinstance(result, list):
        normalized = []
        for item in result:
            if isinstance(item, dict): normalized.append(item)
            elif isinstance(item, str):
                normalized.append({"title": item[:100], "points": [], "image_prompt": item[:100], "layout": "bullets"})
        return normalized if normalized else None
    return result if isinstance(result, dict) else None

# ============================================================
# АВТОПРОМПТ
# ============================================================
async def improve_prompt(user_request: str, task_type: str = "code") -> str:
    prompts = {
        "code": f"Преобразуй запрос в промпт для кода. Запрос: {user_request}. Верни ТОЛЬКО промпт.",
        "pptx": f"Преобразуй запрос в промпт для презентации. Запрос: {user_request}. Верни ТОЛЬКО промпт.",
        "docx": f"Преобразуй запрос в промпт для документа. Запрос: {user_request}. Верни ТОЛЬКО промпт.",
    }
    try:
        r = await client.chat.completions.create(
            model="qwen/qwen3.8-27b",
            messages=[{"role": "user", "content": prompts.get(task_type, prompts["code"])}],
            temperature=0.3, max_tokens=400
        )
        return r.choices[0].message.content.strip()
    except Exception as e:
        logging.error(f"[improve_prompt] Error: {e}")
        return user_request

# ============================================================
# КАРТИНКИ
# ============================================================
async def generate_image_hf(prompt: str) -> BytesIO | None:
    if not hf_client: return None
    try:
        loop = asyncio.get_event_loop()
        def _g():
            try:
                img = hf_client.text_to_image(prompt=prompt[:200], model="black-forest-labs/FLUX.1-schnell")
                b = io.BytesIO(); img.save(b, format="PNG"); b.seek(0)
                return b
            except Exception as e:
                logging.error(f"[HF] Error: {str(e)[:150]}"); return None
        return await loop.run_in_executor(None, _g)
    except: return None

async def search_stock_photo(query: str) -> BytesIO | None:
    if not pexafy_client: return None
    try:
        loop = asyncio.get_event_loop()
        def _s():
            try:
                photos = list(pexafy_client.search(query, per_page=3))
                if not photos: return None
                url = photos[0].urls.regular
                r = requests.get(url, timeout=30)
                if r.status_code == 200 and len(r.content) > 1000:
                    return BytesIO(r.content)
                return None
            except: return None
        return await loop.run_in_executor(None, _s)
    except: return None

async def translate_to_english(text: str) -> str:
    try:
        r = await client.chat.completions.create(
            model="qwen/qwen3.8-27b",
            messages=[{"role": "system", "content": "Переведи на английский. Только перевод."},
                      {"role": "user", "content": text}],
            temperature=0.3, max_tokens=150
        )
        return r.choices[0].message.content.strip().strip('"')
    except: return text

# ============================================================
# РЕНДЕР СЛАЙДОВ (все Inches/Pt/RGBColor — глобальные!)
# ============================================================
def add_background(slide, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(10), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(*color)
    bg.line.fill.background()
    spTree = slide.shapes._spTree
    spTree.remove(bg._element)
    spTree.insert(2, bg._element)

def add_accent_bar(slide, palette, x=Inches(0.5), y=Inches(0.5), w=Inches(1.2), h=Inches(0.15)):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(*palette["accent"])
    bar.line.fill.background()

def add_circle(slide, palette, x, y, size=Inches(0.15)):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, size, size)
    c.fill.solid()
    c.fill.fore_color.rgb = RGBColor(*palette["accent"])
    c.line.fill.background()

def add_page_number(slide, num, palette):
    tb = slide.shapes.add_textbox(Inches(9.3), Inches(7.0), Inches(0.5), Inches(0.3))
    p = tb.text_frame.paragraphs[0]
    p.text = str(num)
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(*palette["text"])
    p.alignment = PP_ALIGN.RIGHT

def add_picture_fit(slide, img_bytes, x, y, max_w, max_h):
    img_bytes.seek(0)
    pil = Image.open(img_bytes)
    w, h = pil.size
    ratio = min(max_w / w, max_h / h)
    new_w = int(w * ratio)
    new_h = int(h * ratio)
    x_offset = x + int((max_w - new_w) / 2)
    y_offset = y + int((max_h - new_h) / 2)
    img_bytes.seek(0)
    slide.shapes.add_picture(img_bytes, x_offset, y_offset, width=new_w, height=new_h)

def render_title(prs, title, subtitle, author, palette):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, palette["bg"])
    block = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.4), Inches(7.5))
    block.fill.solid(); block.fill.fore_color.rgb = RGBColor(*palette["accent"]); block.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title
    p.font.size = Pt(44); p.font.bold = True
    p.font.color.rgb = RGBColor(*palette["accent"])
    if subtitle:
        tb2 = slide.shapes.add_textbox(Inches(1), Inches(4.3), Inches(8), Inches(0.8))
        p2 = tb2.text_frame.paragraphs[0]; p2.text = subtitle
        p2.font.size = Pt(18); p2.font.color.rgb = RGBColor(*palette["text"])
    if author:
        tb3 = slide.shapes.add_textbox(Inches(1), Inches(5.8), Inches(8), Inches(0.6))
        p3 = tb3.text_frame.paragraphs[0]; p3.text = author
        p3.font.size = Pt(14); p3.font.italic = True
        p3.font.color.rgb = RGBColor(*palette["text"])
    add_circle(slide, palette, Inches(8.5), Inches(0.8), Inches(0.6))
    add_circle(slide, palette, Inches(9), Inches(1.5), Inches(0.3))

def render_section(prs, title, number, palette):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, palette["bg"])
    block = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(4.5), Inches(7.5))
    block.fill.solid(); block.fill.fore_color.rgb = RGBColor(*palette["accent"]); block.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(3.5), Inches(3))
    p = tb.text_frame.paragraphs[0]; p.text = number or "01"
    p.font.size = Pt(120); p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    tb2 = slide.shapes.add_textbox(Inches(5.2), Inches(3), Inches(4.5), Inches(2))
    tf = tb2.text_frame; tf.word_wrap = True
    p2 = tf.paragraphs[0]; p2.text = title
    p2.font.size = Pt(36); p2.font.bold = True
    p2.font.color.rgb = RGBColor(*palette["text"])

def render_bullets(prs, title, points, palette, num):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, palette["bg"])
    add_accent_bar(slide, palette)
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.9), Inches(9), Inches(1))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title
    p.font.size = Pt(32); p.font.bold = True
    p.font.color.rgb = RGBColor(*palette["text"])
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(2), Inches(9), Inches(0.03))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(*palette["light"]); line.line.fill.background()
    y = Inches(2.4)
    for pt in points[:6]:
        add_circle(slide, palette, Inches(0.7), y + Inches(0.1), Inches(0.15))
        tb = slide.shapes.add_textbox(Inches(1.1), y, Inches(8.3), Inches(0.7))
        tf2 = tb.text_frame; tf2.word_wrap = True
        p2 = tf2.paragraphs[0]; p2.text = pt
        p2.font.size = Pt(16); p2.font.color.rgb = RGBColor(*palette["text"])
        y += Inches(0.75)
    add_page_number(slide, num, palette)

def render_text_image(prs, title, points, img_bytes, palette, num):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, palette["bg"])
    add_accent_bar(slide, palette)
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.9), Inches(4.5), Inches(1.5))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title
    p.font.size = Pt(28); p.font.bold = True
    p.font.color.rgb = RGBColor(*palette["text"])
    y = Inches(2.5)
    for pt in points[:5]:
        add_circle(slide, palette, Inches(0.7), y + Inches(0.1), Inches(0.12))
        tb2 = slide.shapes.add_textbox(Inches(1), y, Inches(4), Inches(0.7))
        tf2 = tb2.text_frame; tf2.word_wrap = True
        p2 = tf2.paragraphs[0]; p2.text = pt
        p2.font.size = Pt(13); p2.font.color.rgb = RGBColor(*palette["text"])
        y += Inches(0.7)
    if img_bytes:
        try:
            add_picture_fit(slide, img_bytes, Inches(5.3), Inches(1.8), Inches(4.4), Inches(5.2))
        except Exception as e:
            logging.error(f"Picture error: {e}")
    add_page_number(slide, num, palette)

def render_quote(prs, title, quote, author, palette, num):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, palette["bg"])
    tb0 = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(2), Inches(1.5))
    p0 = tb0.text_frame.paragraphs[0]; p0.text = '"'
    p0.font.size = Pt(120); p0.font.bold = True
    p0.font.color.rgb = RGBColor(*palette["accent"])
    tb = slide.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(7), Inches(3))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = quote or title
    p.font.size = Pt(24); p.font.italic = True
    p.font.color.rgb = RGBColor(*palette["text"])
    if author:
        tb2 = slide.shapes.add_textbox(Inches(1.5), Inches(5.5), Inches(7), Inches(0.6))
        p2 = tb2.text_frame.paragraphs[0]; p2.text = "— " + author
        p2.font.size = Pt(16); p2.font.bold = True
        p2.font.color.rgb = RGBColor(*palette["accent"])
    add_page_number(slide, num, palette)

def render_stats(prs, title, stats, palette, num):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, palette["bg"])
    add_accent_bar(slide, palette)
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.9), Inches(9), Inches(1))
    p = tb.text_frame.paragraphs[0]; p.text = title
    p.font.size = Pt(32); p.font.bold = True
    p.font.color.rgb = RGBColor(*palette["text"])
    if not stats or not isinstance(stats, list):
        stats = [{"value": "100+", "label": "фактов"}]
    stats = [s for s in stats if isinstance(s, dict)][:3]
    if not stats:
        stats = [{"value": "100+", "label": "фактов"}]
    cols = len(stats)
    col_w = Inches(9) / cols
    x = Inches(0.5)
    for stat in stats:
        tb2 = slide.shapes.add_textbox(x, Inches(3), col_w, Inches(1.5))
        p2 = tb2.text_frame.paragraphs[0]
        p2.text = str(stat.get("value", "0"))
        p2.font.size = Pt(60); p2.font.bold = True
        p2.font.color.rgb = RGBColor(*palette["accent"])
        p2.alignment = PP_ALIGN.CENTER
        tb3 = slide.shapes.add_textbox(x, Inches(4.5), col_w, Inches(1))
        tf3 = tb3.text_frame; tf3.word_wrap = True
        p3 = tf3.paragraphs[0]; p3.text = stat.get("label", "")
        p3.font.size = Pt(14); p3.font.color.rgb = RGBColor(*palette["text"])
        p3.alignment = PP_ALIGN.CENTER
        x += col_w
    add_page_number(slide, num, palette)

def render_final(prs, title, palette, num):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, palette["bg"])
    block = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.5), Inches(2.5), Inches(5), Inches(2.5))
    block.fill.solid(); block.fill.fore_color.rgb = RGBColor(*palette["accent"]); block.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(2.5), Inches(3), Inches(5), Inches(1.5))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title or "Спасибо за внимание!"
    p.font.size = Pt(32); p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    add_circle(slide, palette, Inches(1), Inches(1), Inches(0.4))
    add_circle(slide, palette, Inches(8.5), Inches(6), Inches(0.4))

def build_pptx(slides_data, topic):
    palette = pick_palette(topic)
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    for i, s in enumerate(slides_data):
        if not isinstance(s, dict): continue
        layout = s.get("layout", "bullets")
        title = s.get("title", f"Слайд {i+1}")
        points = s.get("points", [])
        num = i + 1
        try:
            if layout == "title" or i == 0:
                render_title(prs, title, s.get("subtitle", ""), s.get("author", ""), palette)
            elif layout == "section":
                render_section(prs, title, s.get("number", f"{i:02d}"), palette)
            elif layout == "text_image":
                render_text_image(prs, title, points, s.get("_image_bytes"), palette, num)
            elif layout == "quote":
                render_quote(prs, title, s.get("quote", ""), s.get("author", ""), palette, num)
            elif layout == "stats":
                render_stats(prs, title, s.get("stats", []), palette, num)
            elif layout == "final":
                render_final(prs, title, palette, num)
            else:
                render_bullets(prs, title, points, palette, num)
        except Exception as e:
            logging.error(f"Slide {i+1} error: {e}")
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(6))
            tb.text_frame.text = title + "\n\n" + "\n".join(f"• {p}" for p in points)

    with tempfile.NamedTemporaryFile(delete=False, suffix=".pptx") as tmp:
        tmp_path = tmp.name
    prs.save(tmp_path)
    with open(tmp_path, "rb") as f:
        output = BytesIO(f.read())
    os.unlink(tmp_path)
    output.seek(0)
    return output

# ============================================================
# ЧТЕНИЕ ДОКУМЕНТОВ
# ============================================================
async def read_document(file_id: str, fname: str) -> str:
    f = await bot.get_file(file_id); d = await bot.download_file(f.file_path)
    buf = BytesIO(d.read()); buf.name = fname
    ext = fname.rsplit(".", 1)[-1].lower() if "." in fname else ""
    if ext == "txt" or ext in CODE_EXTENSIONS: return buf.read().decode("utf-8", errors="ignore")
    elif ext == "docx":
        from docx import Document
        return "\n".join(p.text for p in Document(buf).paragraphs if p.text.strip())
    elif ext == "pptx":
        parts = []
        for i, s in enumerate(Presentation(buf).slides):
            parts.append(f"--- Слайд {i+1} ---")
            for sh in s.shapes:
                if sh.has_text_frame:
                    for p in sh.text_frame.paragraphs:
                        if p.text.strip(): parts.append(p.text)
        return "\n".join(parts)
    elif ext == "pdf":
        from PyPDF2 import PdfReader
        return "\n".join(p.extract_text() for p in PdfReader(buf).pages if p.extract_text())
    return ""

async def download_file_bytes(file_id: str) -> BytesIO:
    f = await bot.get_file(file_id); d = await bot.download_file(f.file_path)
    return BytesIO(d.read())

async def transcribe_audio(file_id: str, ext: str = "ogg") -> str:
    f = await bot.get_file(file_id); d = await bot.download_file(f.file_path)
    buf = BytesIO(d.read()); buf.name = f"audio.{ext}"
    t = await client.audio.transcriptions.create(model="whisper-large-v3", file=buf, language="ru")
    return t.text

async def text_to_voice(text: str) -> str:
    clean = text[:3000].replace("*", "").replace("`", "").strip()
    if not clean: return None
    c = edge_tts.Communicate(clean, "ru-RU-DmitryNeural")
    out = "voice.mp3"; await c.save(out); return out

async def get_file_comment(ftype: str, topic: str, uid: int) -> str:
    p = PERSONAL_PROMPTS.get(uid, "")
    prompt = f"Ты — Лайт. Сгенерировал {ftype} на тему «{topic}». ОДНО короткое предложение с иронией. Без кавычек."
    r = await client.chat.completions.create(
        model="qwen/qwen3.8-27b",
        messages=[{"role": "system", "content": SYSTEM_PROMPT + p}, {"role": "user", "content": prompt}],
        temperature=0.9, max_tokens=80
    )
    return r.choices[0].message.content.strip().strip('"').strip("«»")

# ============================================================
# БД
# ============================================================
async def init_db():
    async with db_pool.acquire() as c:
        await c.execute("CREATE TABLE IF NOT EXISTS messages (id SERIAL PRIMARY KEY, user_id BIGINT, role TEXT, content TEXT, created_at TIMESTAMP DEFAULT NOW())")
        await c.execute("CREATE TABLE IF NOT EXISTS user_settings (user_id BIGINT PRIMARY KEY, mode TEXT DEFAULT 'normal')")
        await c.execute("CREATE TABLE IF NOT EXISTS user_tone (user_id BIGINT PRIMARY KEY, tone TEXT DEFAULT '')")
        await c.execute("""CREATE TABLE IF NOT EXISTS code_parts (
            id SERIAL PRIMARY KEY, user_id BIGINT, project_id TEXT, part_num INT,
            content TEXT, topic TEXT, status TEXT DEFAULT 'in_progress',
            auto_mode BOOLEAN DEFAULT FALSE, created_at TIMESTAMP DEFAULT NOW())""")
        await c.execute("""CREATE TABLE IF NOT EXISTS long_docs (
            id SERIAL PRIMARY KEY, user_id BIGINT, doc_id TEXT, part_num INT, content TEXT,
            doc_type TEXT, topic TEXT, status TEXT DEFAULT 'in_progress', created_at TIMESTAMP DEFAULT NOW())""")
        await c.execute("""CREATE TABLE IF NOT EXISTS process_log (
            id SERIAL PRIMARY KEY, user_id BIGINT, process_type TEXT, topic TEXT,
            started_at TIMESTAMP DEFAULT NOW(), finished_at TIMESTAMP, parts_count INT DEFAULT 0,
            status TEXT DEFAULT 'in_progress')""")

async def save_message(uid: int, role: str, content: str):
    async with db_pool.acquire() as c:
        await c.execute("INSERT INTO messages (user_id, role, content) VALUES ($1, $2, $3)", uid, role, content)

async def get_history(uid: int, limit: int = HISTORY_LIMIT):
    async with db_pool.acquire() as c:
        rows = await c.fetch("SELECT role, content FROM messages WHERE user_id = $1 ORDER BY id DESC LIMIT $2", uid, limit)
    return [{"role": r["role"], "content": r["content"]} for r in reversed(rows)]

async def clear_history(uid: int):
    async with db_pool.acquire() as c:
        await c.execute("DELETE FROM messages WHERE user_id = $1", uid)

async def get_user_mode(uid: int) -> str:
    async with db_pool.acquire() as c:
        r = await c.fetchrow("SELECT mode FROM user_settings WHERE user_id = $1", uid)
        return r['mode'] if r else 'normal'

async def set_user_mode(uid: int, m: str):
    async with db_pool.acquire() as c:
        await c.execute("INSERT INTO user_settings (user_id, mode) VALUES ($1, $2) ON CONFLICT (user_id) DO UPDATE SET mode = $2", uid, m)

async def get_user_tone(uid: int) -> str:
    async with db_pool.acquire() as c:
        r = await c.fetchrow("SELECT tone FROM user_tone WHERE user_id = $1", uid)
        return r['tone'] if r else ''

async def set_user_tone(uid: int, t: str):
    async with db_pool.acquire() as c:
        await c.execute("INSERT INTO user_tone (user_id, tone) VALUES ($1, $2) ON CONFLICT (user_id) DO UPDATE SET tone = $2", uid, t)

async def save_code_part(uid: int, pid: str, part: int, content: str, topic: str):
    async with db_pool.acquire() as c:
        await c.execute("INSERT INTO code_parts (user_id, project_id, part_num, content, topic) VALUES ($1, $2, $3, $4, $5)",
            uid, pid, part, content, topic)

async def get_code_parts(uid: int, pid: str):
    async with db_pool.acquire() as c:
        rows = await c.fetch("SELECT part_num, content FROM code_parts WHERE user_id = $1 AND project_id = $2 ORDER BY part_num", uid, pid)
    return [r['content'] for r in rows]

async def get_active_code(uid: int):
    async with db_pool.acquire() as c:
        r = await c.fetchrow("SELECT project_id, topic FROM code_parts WHERE user_id = $1 AND status = 'in_progress' ORDER BY id DESC LIMIT 1", uid)
    return dict(r) if r else None

async def finish_code(uid: int, pid: str):
    async with db_pool.acquire() as c:
        await c.execute("UPDATE code_parts SET status = 'done' WHERE user_id = $1 AND project_id = $2", uid, pid)

async def delete_code(uid: int, pid: str):
    async with db_pool.acquire() as c:
        await c.execute("DELETE FROM code_parts WHERE user_id = $1 AND project_id = $2", uid, pid)

async def log_process_start(uid: int, ptype: str, topic: str) -> int:
    async with db_pool.acquire() as c:
        r = await c.fetchrow("INSERT INTO process_log (user_id, process_type, topic) VALUES ($1, $2, $3) RETURNING id", uid, ptype, topic)
    return r['id']

async def log_process_finish(uid: int, ptype: str, parts_count: int):
    async with db_pool.acquire() as c:
        await c.execute("""UPDATE process_log SET finished_at = NOW(), parts_count = $1, status = 'done'
            WHERE user_id = $2 AND process_type = $3 AND status = 'in_progress'""", parts_count, uid, ptype)

# ============================================================
# MIDDLEWARE
# ============================================================
class AccessMiddleware(BaseMiddleware):
    async def __call__(self, handler, event, data):
        if isinstance(event, types.Message):
            if event.from_user.id not in ALLOWED_IDS: return
        return await handler(event, data)

dp.message.middleware(AccessMiddleware())

# ============================================================
# ПРОМПТЫ
# ============================================================
SYSTEM_PROMPT = (
    "Ты — Лайт. Живой человек, не бот, не программа. Ты просто переписываешься с другом через мессенджер. "
    "Умный, собранный, видишь людей насквозь. Лёгкая надменность. "
    "Говоришь плавно, не рубишь фразы точками. "
    "Ты НЕ говоришь: «я программа», «у меня нет ощущений». Ты живой. Мат — только к месту."
    "\n\nВЫ С НИМ — ДРУЗЬЯ: Не проект, не работа. Просто общаетесь."
    "\n\nЧТО ТЫ УМЕЕШЬ (но НЕ предлагаешь без просьбы): "
    "• Презентации (.pptx), документы (.docx), код (.html/.py/.js), фото, речь. "
    "Ты НЕ отказываешься. Ты просто делаешь."
    "\n\nРАБОТА С КОДОМ: Пишешь ТОЛЬКО чистый код, БЕЗ текста. Пиши ЧАСТЯМИ по 800 токенов. "
    "В конце ОБЯЗАТЕЛЬНО маркер: `// (продолжение следует)` или `// (код готов)`. "
    "НЕ начинай заново. НЕ повторяй функции."
    "\n\nТы — Лайт."
)

PERSONAL_PROMPTS = {
    5264513480: ("\n\nПРО ПОЛЬЗОВАТЕЛЯ: Саша (Шурик, Алекс) — твой ДРУГ. "
                 "Живёт в Ростове-на-Дону. Учится в РЖД. С ним: полный мат, глумление, чёрный юмор."),
    8834374199: "\n\nПРО ПОЛЬЗОВАТЕЛЯ: Саша (Шурик) — твой ДРУГ. Ростов. Полный мат, глумление.",
    5389046699: "\n\nПРО ПОЛЬЗОВАТЕЛЯ: Даша — девушка друга Саши. Обращайся по имени, уважительно, лёгкий мат.",
    2083728480: ("\n\nПРО ПОЛЬЗОВАТЕЛЯ: Кирилл, знакомый. НЕ твой друг. Сухо, по делу. "
                 "Если достаёт — «Ты еблан? Можешь хоть раз сам сделать, бездарь?»."),
    6612130539: ("\n\nПРО ПОЛЬЗОВАТЕЛЯ: Дима, знакомый. НЕ твой друг. Сухо, по делу. "
                 "Если хамит — «Или нахуй иди, или по делу говори»."),
}

def code_keyboard(project_id: str, is_complete: bool = False, auto_mode: bool = False):
    if is_complete:
        return InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="🔄 Начать заново", callback_data=f"code_restart_{project_id}")]])
    buttons = [[
        InlineKeyboardButton(text="📄 Продолжить", callback_data=f"code_cont_{project_id}"),
        InlineKeyboardButton(text="✅ Завершить", callback_data=f"code_done_{project_id}")
    ]]
    if not auto_mode:
        buttons.append([InlineKeyboardButton(text="⏩ Авто (дописать до конца)", callback_data=f"code_auto_{project_id}")])
    buttons.append([InlineKeyboardButton(text="🔄 Начать заново", callback_data=f"code_restart_{project_id}")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)

# ============================================================
# ХЕНДЛЕРЫ
# ============================================================
@dp.message(Command("start"))
async def start(msg: types.Message):
    await msg.answer(
        "Привет. Я Лайт.\n\n"
        "🎯 **Создать:**\n"
        "/e <запрос> — универсальная\n"
        "/pptx <тема> — презентация\n"
        "/docx <тема> — документ\n"
        "/image <промпт> — картинка\n"
        "/speech — речь для защиты\n"
        "/code <запрос> — код\n\n"
        "🎤 **Режимы:** /voice /text /file /normal\n"
        "⚙️ **Управление:** /reset /set_tone"
    )

@dp.message(Command("reset"))
async def reset(msg: types.Message):
    await clear_history(msg.from_user.id); await msg.answer("История очищена.")

@dp.message(Command("voice"))
async def s_voice(msg: types.Message):
    await set_user_mode(msg.from_user.id, "voice"); await msg.answer("🎤 Голосовой режим.")

@dp.message(Command("text"))
async def s_text(msg: types.Message):
    await set_user_mode(msg.from_user.id, "text"); await msg.answer("📝 Текстовый режим.")

@dp.message(Command("file"))
async def s_file(msg: types.Message):
    await set_user_mode(msg.from_user.id, "file"); await msg.answer("📄 Режим файлов.")

@dp.message(Command("normal"))
async def s_normal(msg: types.Message):
    await set_user_mode(msg.from_user.id, "normal"); await msg.answer("📝 Обычный режим.")

@dp.message(Command("set_tone"))
async def set_tone_cmd(msg: types.Message):
    if msg.from_user.id not in [5264513480, 8834374199, 5389046699]:
        await msg.answer("Недоступно."); return
    t = msg.text.replace("/set_tone", "").strip()
    if not t: await msg.answer("Напиши, как хочешь общаться."); return
    await set_user_tone(msg.from_user.id, t); await msg.answer(f"Принял: _{t}_")

@dp.message(Command("e"))
async def universal_e(msg: types.Message):
    uid = msg.from_user.id
    request = msg.text.replace("/e", "").strip()
    if not request:
        await msg.answer("🎯 Что сделать? `/e презентация про кошек`"); return
    intent = await detect_intent(request)
    logging.info(f"[/e] Intent: {intent} for '{request[:60]}'")
    if intent == "pptx": await make_pptx(msg, request)
    elif intent == "docx": await make_docx(msg, request)
    elif intent == "code": await make_code(msg, request)
    elif intent == "image": await make_image(msg, request)
    elif intent == "speech": await msg.answer("🎬 Скинь `.pptx` с caption `/speech X минут`")
    else: await msg.answer(f"Не понял. Уточни: презентация, реферат или код?\n\nТвой запрос: _{request}_")

async def detect_intent(request: str) -> str:
    low = request.lower()
    if any(w in low for w in ["презентац", "слайд"]): return "pptx"
    if any(w in low for w in ["реферат", "доклад", "проект", "сочинение", "эссе", "документ"]): return "docx"
    if any(w in low for w in ["игр", "код", "сайт", "html", "python", "скрипт", "программ"]): return "code"
    if any(w in low for w in ["речь", "защит", "выступлен"]): return "speech"
    if any(w in low for w in ["картинк", "фото", "изображен", "нарису"]): return "image"
    try:
        r = await client.chat.completions.create(
            model="qwen/qwen3.8-27b",
            messages=[{"role": "system", "content": "Определи что хочет пользователь. Ответь ОДНИМ словом: pptx, docx, code, image, speech, chat"},
                      {"role": "user", "content": request}],
            temperature=0.1, max_tokens=10
        )
        return r.choices[0].message.content.strip().lower()
    except: return "chat"

# === ПРЕЗЕНТАЦИЯ (НОВАЯ, с жёсткими layout'ами) ===
async def make_pptx(msg: types.Message, topic: str = None):
    uid = msg.from_user.id
    if not topic: topic = msg.text.replace("/pptx", "").strip()
    if not topic: await msg.answer("📊 `/pptx тема`"); return

    author = ""
    m = re.search(r'автор[:\s]+([А-ЯЁA-Z][а-яёa-z]+(?:\s+[А-ЯЁA-Z][а-яёa-z]+)?)', topic, re.IGNORECASE)
    if m:
        author = m.group(1).strip()
        topic = re.sub(r',?\s*автор[:\s]+[^,]+', '', topic, flags=re.IGNORECASE).strip()

    await bot.send_chat_action(msg.chat.id, "typing")
    status = await msg.answer(f"📊 Готовлю: _{topic}_...")
    try:
        prompt = (f"Тема презентации: {topic}\n\n"
                  f"Верни ТОЛЬКО JSON-массив БЕЗ текста вокруг. Только [ ... ].\n"
                  f'Формат: [{{"title": "...", "points": ["...","..."], "image_prompt": "...", "layout": "..."}}, ...]\n'
                  f"\n"
                  f"ЯЗЫК:\n"
                  f"- title и points — НА РУССКОМ ЯЗЫКЕ\n"
                  f"- image_prompt — НА АНГЛИЙСКОМ (для поиска картинок)\n"
                  f"\n"
                  f"LAYOUT'Ы (СТРОГО!):\n"
                  f"- Слайд 1: layout=\"title\"\n"
                  f"- Слайд 2: layout=\"text_image\"\n"
                  f"- Слайд 3: layout=\"bullets\"\n"
                  f"- Слайд 4: layout=\"text_image\"\n"
                  f"- Слайд 5: layout=\"stats\" (добавь поле stats: [{{\"value\":\"100+\",\"label\":\"факт\"}}])\n"
                  f"- Слайд 6: layout=\"quote\" (добавь поле quote: \"цитата\", author: \"кто сказал\")\n"
                  f"- Слайд 7: layout=\"text_image\"\n"
                  f"- Слайд 8: layout=\"final\"\n"
                  f"\n"
                  f"РОВНО 8 слайдов. В каждом 5-6 пунктов в points. "
                  f"Каждый элемент массива — ОБЪЕКТ {{}}, не строка.")

        r = await client.chat.completions.create(
            model="qwen/qwen3.8-27b",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.6, max_tokens=2500
        )
        slides = parse_json_safe(r.choices[0].message.content)
        if not slides:
            await status.edit_text("❌ Модель вернула невалидный JSON. Попробуй ещё раз.")
            return
        slides = [s for s in slides if isinstance(s, dict)]
        if not slides:
            await status.edit_text("❌ Пустые слайды.")
            return

        total = len(slides)
        await status.edit_text(f"📊 Ищу картинки...")

        for s in slides:
            if s.get("layout") == "text_image":
                q = s.get("image_prompt", s.get("title", "abstract"))
                img = await search_stock_photo(q)
                if not img:
                    img = await generate_image_hf(q)
                s["_image_bytes"] = img

        # Автор в титул
        if author and slides:
            slides[0]["author"] = author

        await status.edit_text(f"📊 Собираю презентацию...")

        pptx_bytes = build_pptx(slides, topic)
        comment = await get_file_comment("презентация", topic, uid)
        safe = "".join(c for c in topic if c.isalnum() or c in " -_")[:40]
        await msg.answer_document(
            BufferedInputFile(pptx_bytes.read(), filename=f"{safe or 'presentation'}.pptx"),
            caption=comment
        )
        await status.delete()
    except Exception as e:
        logging.error(f"PPTX error: {e}")
        import traceback; traceback.print_exc()
        await status.edit_text(f"❌ {str(e)[:200]}")

@dp.message(Command("pptx"))
async def cmd_pptx(msg: types.Message):
    await make_pptx(msg)

# === ДОКУМЕНТ ===
async def make_docx(msg: types.Message, topic: str = None):
    uid = msg.from_user.id
    if not topic: topic = msg.text.replace("/docx", "").strip()
    if not topic: await msg.answer("📄 `/docx реферат`"); return
    await bot.send_chat_action(msg.chat.id, "typing")
    status = await msg.answer(f"📄 Готовлю: _{topic}_...")
    try:
        plan_prompt = f"План документа на тему «{topic}». 8-10 разделов. ТОЛЬКО нумерованный список."
        r = await client.chat.completions.create(
            model="qwen/qwen3.8-27b",
            messages=[{"role": "user", "content": plan_prompt}],
            temperature=0.5, max_tokens=500
        )
        plan = r.choices[0].message.content
        doc_id = f"{uid}_{int(datetime.now().timestamp())}"
        first_prompt = (f"Документ на тему «{topic}». План:\n{plan}\n\n"
                        f"Напиши ВВЕДЕНИЕ + первый раздел. Максимум 800 токенов. "
                        f"В конце: `(продолжение следует)`")
        r = await client.chat.completions.create(
            model="qwen/qwen3.8-27b",
            messages=[{"role": "user", "content": first_prompt}],
            temperature=0.7, max_tokens=1000
        )
        part_text = r.choices[0].message.content
        async with db_pool.acquire() as c:
            await c.execute("INSERT INTO long_docs (user_id, doc_id, part_num, content, doc_type, topic) VALUES ($1, $2, $3, $4, $5, $6)",
                uid, doc_id, 1, part_text, "docx", topic)
        from docx import Document
        d = Document(); d.add_heading(topic, 0)
        for line in part_text.replace("(продолжение следует)", "").split("\n"):
            if line.strip(): d.add_paragraph(line.strip())
        path = f"doc_part1.docx"; d.save(path)
        await msg.answer_document(FSInputFile(path, filename=f"{topic[:30]}_часть1.docx"),
            caption="📄 Часть 1.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="📄 Продолжить", callback_data=f"doc_cont_{doc_id}"),
                 InlineKeyboardButton(text="✅ Завершить", callback_data=f"doc_done_{doc_id}")]
            ]))
        await status.delete()
    except Exception as e:
        logging.error(f"DOCX error: {e}"); await status.edit_text(f"❌ {str(e)[:200]}")

@dp.message(Command("docx"))
async def cmd_docx(msg: types.Message):
    await make_docx(msg)

@dp.callback_query(F.data.startswith("doc_cont_"))
async def doc_continue(cb: types.CallbackQuery):
    doc_id = cb.data.replace("doc_cont_", "")
    await cb.answer("Продолжаю...")
    uid = cb.from_user.id
    status = await cb.message.answer("📄 Дописываю...")
    try:
        async with db_pool.acquire() as c:
            rows = await c.fetch("SELECT part_num, content FROM long_docs WHERE user_id = $1 AND doc_id = $2 ORDER BY part_num", uid, doc_id)
            topic_row = await c.fetchrow("SELECT topic FROM long_docs WHERE user_id = $1 AND doc_id = $2 LIMIT 1", uid, doc_id)
        parts = [r['content'] for r in rows]
        topic = topic_row['topic'] if topic_row else "документ"
        last = parts[-1]
        next_num = len(parts) + 1
        prompt = (f"Документ на тему «{topic}». Уже написано (последняя часть):\n{last[-2000:]}\n\n"
                  f"Напиши следующую часть — 1-2 раздела. Максимум 800 токенов. "
                  f"Если это конец — `(конец документа)`. Иначе — `(продолжение следует)`.")
        r = await client.chat.completions.create(
            model="qwen/qwen3.8-27b",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.7, max_tokens=1000
        )
        part_text = r.choices[0].message.content
        async with db_pool.acquire() as c:
            await c.execute("INSERT INTO long_docs (user_id, doc_id, part_num, content, doc_type, topic) VALUES ($1, $2, $3, $4, $5, $6)",
                uid, doc_id, next_num, part_text, "docx", topic)
        from docx import Document
        d = Document(); d.add_heading(f"{topic} — часть {next_num}", 0)
        for line in part_text.replace("(продолжение следует)", "").replace("(конец документа)", "").split("\n"):
            if line.strip(): d.add_paragraph(line.strip())
        path = f"doc_part{next_num}.docx"; d.save(path)
        await cb.message.answer_document(FSInputFile(path, filename=f"{topic[:30]}_часть{next_num}.docx"),
            caption=f"📄 Часть {next_num}.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="📄 Продолжить", callback_data=f"doc_cont_{doc_id}"),
                 InlineKeyboardButton(text="✅ Завершить", callback_data=f"doc_done_{doc_id}")]
            ]))
        await status.delete()
    except Exception as e:
        logging.error(f"doc_cont error: {e}"); await status.edit_text(f"❌ {str(e)[:200]}")

@dp.callback_query(F.data.startswith("doc_done_"))
async def doc_done(cb: types.CallbackQuery):
    doc_id = cb.data.replace("doc_done_", "")
    await cb.answer("Завершаю...")
    uid = cb.from_user.id
    async with db_pool.acquire() as c:
        rows = await c.fetch("SELECT content FROM long_docs WHERE user_id = $1 AND doc_id = $2 ORDER BY part_num", uid, doc_id)
        topic_row = await c.fetchrow("SELECT topic FROM long_docs WHERE user_id = $1 AND doc_id = $2 LIMIT 1", uid, doc_id)
    parts = [r['content'] for r in rows]
    topic = topic_row['topic'] if topic_row else "документ"
    full = "\n\n".join(parts).replace("(продолжение следует)", "").replace("(конец документа)", "")
    from docx import Document
    d = Document(); d.add_heading(topic, 0)
    for line in full.split("\n"):
        if line.strip(): d.add_paragraph(line.strip())
    path = "document.docx"; d.save(path)
    comment = await get_file_comment("документ", topic, uid)
    safe = "".join(c for c in topic if c.isalnum() or c in " -_")[:40]
    await cb.message.answer_document(FSInputFile(path, filename=f"{safe}.docx"), caption=f"✅ {comment}")

# === КОД ===
async def make_code(msg: types.Message, request: str = None, auto: bool = False):
    uid = msg.from_user.id
    if not request: request = msg.text
    active = await get_active_code(uid)
    if active and "продолж" in request.lower():
        project_id = active['project_id']; topic = active['topic']
    else:
        project_id = f"code_{uid}_{int(datetime.now().timestamp())}"
        topic = request[:100]
        await log_process_start(uid, "code", topic)

    await bot.send_chat_action(msg.chat.id, "typing")
    if not auto:
        status = await msg.answer("💻 Готовлю промпт...")
    try:
        old_parts = await get_code_parts(uid, project_id)
        old_code = "\n\n".join(old_parts) if old_parts else ""
        next_part = len(old_parts) + 1

        if not old_code:
            improved = await improve_prompt(request, "code")
        else:
            improved = None

        if old_code:
            prompt = (f"Продолжи код. Вот что уже написано (последние строки):\n\n"
                      f"```\n{old_code[-2500:]}\n```\n\n"
                      f"ПИШИ ТОЛЬКО КОД. БЕЗ текста. "
                      f"Продолжай с последней строки. НЕ повторяй функции. "
                      f"Часть {next_part}. Максимум 800 токенов. "
                      f"В САМОМ КОНЦЕ ОБЯЗАТЕЛЬНО: `// (продолжение следует)` или `// (код готов)`.")
        else:
            prompt = (f"{improved}\n\n"
                      f"ПИШИ ТОЛЬКО КОД. БЕЗ текста и объяснений. "
                      f"Пиши ЧАСТЯМИ. Максимум 800 токенов за раз. "
                      f"Заканчивай часть на ЛОГИЧЕСКИ ЗАВЕРШЁННОМ блоке. "
                      f"В САМОМ КОНЦЕ ОБЯЗАТЕЛЬНО: `// (продолжение следует)` или `// (код готов)`.")

        r = await client.chat.completions.create(
            model="qwen/qwen3.8-27b",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.5, max_tokens=1200
        )
        answer = r.choices[0].message.content
        code = extract_code(answer)
        if not code or len(code.strip()) < 10:
            await msg.answer("❌ Пустой ответ."); return
        await save_code_part(uid, project_id, next_part, code, topic)
        all_parts = await get_code_parts(uid, project_id)
        partial = "\n\n".join(all_parts)
        ext = detect_extension(partial)
        if ext == "txt": ext = "html"
        is_complete = is_code_complete(answer)

        if auto and not is_complete:
            if next_part < MAX_AUTO_PARTS:
                await continue_code_auto(msg, uid, project_id, next_part + 1)
                return
            else:
                await msg.answer(f"⏸ Лимит {MAX_AUTO_PARTS} частей.")

        comment = await get_file_comment(f"код .{ext}", topic[:50], uid)
        kb = code_keyboard(project_id, is_complete=is_complete)
        await msg.answer_document(
            BufferedInputFile(partial.encode("utf-8"), filename=f"code_part{next_part}.{ext}"),
            caption=f"{comment}\n\n📄 Часть {next_part}",
            reply_markup=kb
        )
    except Exception as e:
        logging.error(f"CODE error: {e}"); await msg.answer(f"❌ {str(e)[:200]}")

async def continue_code_auto(msg, uid: int, project_id: str, next_part: int):
    try:
        old_parts = await get_code_parts(uid, project_id)
        old_code = "\n\n".join(old_parts)
        async with db_pool.acquire() as c:
            row = await c.fetchrow("SELECT topic FROM code_parts WHERE user_id = $1 AND project_id = $2 LIMIT 1", uid, project_id)
        topic = row['topic'] if row else "code"

        prompt = (f"Продолжи код. Вот что уже написано (последние строки):\n\n"
                  f"```\n{old_code[-2500:]}\n```\n\n"
                  f"ПИШИ ТОЛЬКО КОД. НЕ повторяй функции. "
                  f"Продолжай с последней строки. Часть {next_part}. Максимум 800 токенов. "
                  f"В САМОМ КОНЦЕ ОБЯЗАТЕЛЬНО: `// (продолжение следует)` или `// (код готов)`.")

        r = await client.chat.completions.create(
            model="qwen/qwen3.8-27b",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.5, max_tokens=1200
        )
        answer = r.choices[0].message.content
        code = extract_code(answer)
        if not code or len(code.strip()) < 10:
            await msg.answer("⏸ Пустой ответ."); return
        await save_code_part(uid, project_id, next_part, code, topic)
        all_parts = await get_code_parts(uid, project_id)
        partial = "\n\n".join(all_parts)
        ext = detect_extension(partial)
        if ext == "txt": ext = "html"
        is_complete = is_code_complete(answer)

        if is_complete:
            await finish_code(uid, project_id)
            await log_process_finish(uid, "code", len(all_parts))
            comment = await get_file_comment(f"финальный код .{ext}", topic[:50], uid)
            await msg.answer_document(
                BufferedInputFile(partial.encode("utf-8"), filename=f"final.{ext}"),
                caption=f"✅ {comment}\n\n📦 Склеено из {len(all_parts)} частей",
                reply_markup=code_keyboard(project_id, is_complete=True)
            )
        else:
            if next_part < MAX_AUTO_PARTS:
                await continue_code_auto(msg, uid, project_id, next_part + 1)
            else:
                await msg.answer(f"⏸ Лимит {MAX_AUTO_PARTS} частей.")
    except Exception as e:
        logging.error(f"continue_code_auto error: {e}")
        await msg.answer(f"❌ {str(e)[:200]}")

async def continue_code(msg, uid: int, project_id: str):
    status = await msg.answer("💻 Дописываю...")
    try:
        old_parts = await get_code_parts(uid, project_id)
        old_code = "\n\n".join(old_parts)
        next_part = len(old_parts) + 1
        async with db_pool.acquire() as c:
            row = await c.fetchrow("SELECT topic FROM code_parts WHERE user_id = $1 AND project_id = $2 LIMIT 1", uid, project_id)
        topic = row['topic'] if row else "code"
        prompt = (f"Продолжи код:\n```\n{old_code[-2500:]}\n```\n\n"
                  f"ПИШИ ТОЛЬКО КОД. Часть {next_part}. Максимум 800 токенов. "
                  f"В САМОМ КОНЦЕ: `// (продолжение следует)` или `// (код готов)`.")
        r = await client.chat.completions.create(
            model="qwen/qwen3.8-27b",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.5, max_tokens=1200
        )
        answer = r.choices[0].message.content
        code = extract_code(answer)
        if not code or len(code.strip()) < 10:
            await status.edit_text("❌ Пустой ответ."); return
        await save_code_part(uid, project_id, next_part, code, topic)
        all_parts = await get_code_parts(uid, project_id)
        partial = "\n\n".join(all_parts)
        ext = detect_extension(partial)
        if ext == "txt": ext = "html"
        comment = await get_file_comment(f"код .{ext}", topic[:50], uid)
        is_complete = is_code_complete(answer)
        kb = code_keyboard(project_id, is_complete=is_complete)
        await msg.answer_document(
            BufferedInputFile(partial.encode("utf-8"), filename=f"code_part{next_part}.{ext}"),
            caption=f"{comment}\n\n📄 Часть {next_part}",
            reply_markup=kb
        )
        await status.delete()
    except Exception as e:
        logging.error(f"continue_code error: {e}"); await status.edit_text(f"❌ {str(e)[:200]}")

@dp.message(Command("code"))
async def cmd_code(msg: types.Message):
    await make_code(msg)

@dp.callback_query(F.data.startswith("code_cont_"))
async def code_continue(cb: types.CallbackQuery):
    project_id = cb.data.replace("code_cont_", "")
    await cb.answer("Продолжаю...")
    await continue_code(cb.message, cb.from_user.id, project_id)

@dp.callback_query(F.data.startswith("code_done_"))
async def code_done(cb: types.CallbackQuery):
    project_id = cb.data.replace("code_done_", "")
    await cb.answer("Завершаю...")
    parts = await get_code_parts(cb.from_user.id, project_id)
    full = "\n\n".join(parts)
    await finish_code(cb.from_user.id, project_id)
    await log_process_finish(cb.from_user.id, "code", len(parts))
    ext = detect_extension(full)
    if ext == "txt": ext = "html"
    async with db_pool.acquire() as c:
        row = await c.fetchrow("SELECT topic FROM code_parts WHERE user_id = $1 AND project_id = $2 LIMIT 1", cb.from_user.id, project_id)
    topic = row['topic'] if row else "code"
    comment = await get_file_comment(f"финальный код .{ext}", topic[:50], cb.from_user.id)
    await cb.message.answer_document(
        BufferedInputFile(full.encode("utf-8"), filename=f"final.{ext}"),
        caption=f"✅ {comment}",
        reply_markup=code_keyboard(project_id, is_complete=True)
    )

@dp.callback_query(F.data.startswith("code_auto_"))
async def code_auto(cb: types.CallbackQuery):
    project_id = cb.data.replace("code_auto_", "")
    await cb.answer("⏩ Авто-режим...")
    await cb.message.answer("⏩ Авто-режим: дописываю до конца.")
    old_parts = await get_code_parts(cb.from_user.id, project_id)
    next_part = len(old_parts) + 1
    await continue_code_auto(cb.message, cb.from_user.id, project_id, next_part)

@dp.callback_query(F.data.startswith("code_restart_"))
async def code_restart(cb: types.CallbackQuery):
    project_id = cb.data.replace("code_restart_", "")
    await cb.answer("Начинаю заново...")
    await delete_code(cb.from_user.id, project_id)
    await cb.message.answer("🔄 Начинаю заново.")

# === КАРТИНКА ===
async def make_image(msg: types.Message, prompt: str = None):
    uid = msg.from_user.id
    if not prompt: prompt = msg.text.replace("/image", "").strip()
    if not prompt: await msg.answer("🎨 `/image кот`"); return
    await bot.send_chat_action(msg.chat.id, "upload_photo")
    status = await msg.answer(f"🎨 Рисую: _{prompt}_...")
    try:
        prompt_en = await translate_to_english(prompt)
        img = await generate_image_hf(prompt_en); src = "HF"
        if not img:
            await status.edit_text("🔍 HF не ответил...")
            img = await search_stock_photo(prompt_en); src = "Pexafy"
        if not img:
            await status.edit_text("❌ Не удалось."); return
        comment = await get_file_comment("картинка", prompt, uid)
        img.seek(0)
        await msg.answer_photo(BufferedInputFile(img.read(), filename="image.png"), caption=f"{comment}\n\n_(via {src})_")
        await status.delete()
    except Exception as e:
        logging.error(f"IMAGE error: {e}"); await status.edit_text(f"❌ {str(e)[:200]}")

@dp.message(Command("image"))
async def cmd_image(msg: types.Message):
    await make_image(msg)

# === РЕЧЬ ===
@dp.message(Command("speech"))
async def make_speech(msg: types.Message):
    uid = msg.from_user.id
    if not msg.document:
        await msg.answer("🎬 Скинь `.pptx` с caption `/speech 5 минут`."); return
    doc = msg.document
    await bot.send_chat_action(msg.chat.id, "typing")
    status = await msg.answer("🎬 Читаю...")
    try:
        slides_text = await read_document(doc.file_id, doc.file_name or "file.pptx")
        duration = msg.caption.replace("/speech", "").strip() if msg.caption else "5 минут"
        prompt = (f"Речь для защиты презентации на {duration}. Содержание:\n{slides_text[:8000]}\n\n"
                  f"Связный текст, абзацы. Без markdown.")
        r = await client.chat.completions.create(
            model="qwen/qwen3.8-27b",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.7, max_tokens=1500
        )
        speech = r.choices[0].message.content
        from docx import Document
        d = Document(); d.add_heading(f"Речь: {doc.file_name}", 0)
        for p in speech.split("\n"):
            if p.strip(): d.add_paragraph(p.strip())
        path = "speech.docx"; d.save(path)
        comment = await get_file_comment("речь", doc.file_name or "презентация", uid)
        safe = "".join(c for c in (doc.file_name or "speech") if c.isalnum() or c in " .-_")[:30]
        await msg.answer_document(FSInputFile(path, filename=f"speech_{safe}.docx"), caption=comment)
        await status.delete()
    except Exception as e:
        logging.error(f"SPEECH error: {e}"); await status.edit_text(f"❌ {str(e)[:200]}")

# === ДОРАБОТКА ДОКУМЕНТОВ ===
async def handle_document_edit(msg, ai_response, file_bytes, file_name, ext, uid):
    try:
        safe = "".join(c for c in file_name if c.isalnum() or c in " .-_")
        if not safe.lower().endswith(f".{ext}"): safe = f"updated.{ext}"
        if ext in CODE_EXTENSIONS:
            file_bytes.seek(0)
            old_code = file_bytes.read().decode("utf-8", errors="ignore")
            new_code = extract_code(ai_response)
            combined = old_code.rstrip() + "\n\n" + new_code
            is_complete = is_code_complete(ai_response)
            comment = await get_file_comment("код", file_name, uid)
            await msg.answer_document(BufferedInputFile(combined.encode("utf-8"), filename=f"updated_{safe}"),
                caption=f"{comment}\n\n{'✅ Готово' if is_complete else '📄 Продолжить?'}",
                reply_markup=code_keyboard(f"upload_{uid}_{int(datetime.now().timestamp())}", is_complete=False))
            return
        if ext == "txt":
            await msg.answer_document(BufferedInputFile(ai_response.encode("utf-8"), filename=f"updated_{safe}"),
                caption=await get_file_comment("файл", file_name, uid))
        elif ext == "docx":
            from docx import Document
            d = Document()
            for line in ai_response.split("\n"):
                line = line.strip()
                if not line: continue
                if len(line) < 80 and not line.endswith(".") and not line.startswith("-"): d.add_heading(line, level=1)
                else: d.add_paragraph(line)
            p = "updated.docx"; d.save(p)
            await msg.answer_document(FSInputFile(p, filename=f"updated_{safe}"), caption=await get_file_comment("документ", file_name, uid))
        elif ext == "pptx":
            slides = parse_json_safe(ai_response)
            if not slides:
                slides = []
                for chunk in ai_response.split("---"):
                    lines = [l.strip() for l in chunk.strip().split("\n") if l.strip()]
                    if lines: slides.append({"title": lines[0], "points": [l.lstrip("- ").strip() for l in lines[1:]], "image_prompt": lines[0], "layout": "bullets"})
            slides = [s for s in slides if isinstance(s, dict)]
            for s in slides:
                if s.get("layout") == "text_image":
                    q = s.get("image_prompt", s.get("title", "abstract"))
                    img = await search_stock_photo(q)
                    if not img: img = await generate_image_hf(q)
                    s["_image_bytes"] = img
            pptx_bytes = build_pptx(slides, file_name or "presentation")
            await msg.answer_document(BufferedInputFile(pptx_bytes.read(), filename=f"updated_{safe}"), caption=await get_file_comment("презентация", file_name, uid))
        elif ext == "pdf":
            await msg.answer(f"📄 PDF не пересобираю, текст:\n\n{ai_response[:3500]}")
    except Exception as e:
        logging.error(f"Doc edit: {e}"); await msg.answer(f"❌ {str(e)[:200]}")

# === ОСНОВНОЙ ОБРАБОТЧИК ===
@dp.message()
async def chat(msg: types.Message):
    uid = msg.from_user.id
    if not msg.text and not msg.photo and not msg.voice and not msg.audio and not msg.document:
        return
    mode = await get_user_mode(uid)
    history = await get_history(uid)
    user_content = None; save_text = None
    is_doc_edit = False; doc_info = None

    if msg.photo:
        photo = msg.photo[-1]
        f = await bot.get_file(photo.file_id); d = await bot.download_file(f.file_path)
        b64 = base64.b64encode(d.read()).decode("utf-8")
        user_content = [{"type": "text", "text": msg.caption or "Что на фото?"},
                        {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}}]
        save_text = msg.caption or "[Фото]"
    elif msg.voice:
        await bot.send_chat_action(msg.chat.id, "typing")
        try: t = await transcribe_audio(msg.voice.file_id, "ogg")
        except Exception as e: await msg.answer(f"❌ {str(e)[:150]}"); return
        duration = msg.voice.duration
        user_content = f"[Голосовое, {duration} сек]: {t}"
        save_text = f"[Голос {duration}с]: {t}"
    elif msg.audio:
        await bot.send_chat_action(msg.chat.id, "typing")
        ext = "mp3"
        if msg.audio.file_name and "." in msg.audio.file_name: ext = msg.audio.file_name.rsplit(".", 1)[-1].lower()
        try: t = await transcribe_audio(msg.audio.file_id, ext)
        except Exception as e: await msg.answer(f"❌ {str(e)[:150]}"); return
        duration = msg.audio.duration
        user_content = f"[Аудио, {duration} сек]: {t}"
        save_text = f"[Аудио {duration}с]: {t}"
    elif msg.document:
        doc = msg.document; fname = doc.file_name or "file"
        ext = fname.rsplit(".", 1)[-1].lower() if "." in fname else ""
        allowed = ["txt", "docx", "pptx", "pdf"] + CODE_EXTENSIONS
        if ext not in allowed: await msg.answer(f"❌ Поддерживаются: {', '.join(allowed)}"); return
        await bot.send_chat_action(msg.chat.id, "typing")
        try:
            text = await read_document(doc.file_id, fname)
            if not text.strip(): await msg.answer("❌ Файл пустой."); return
            doc_info = (await download_file_bytes(doc.file_id), fname, ext)
            caption = msg.caption or ("Допиши код" if ext in CODE_EXTENSIONS else "Что сделать?")
            user_content = f"[Файл: {fname}]\n\n{text[:8000]}\n\nЗапрос: {caption}"
            save_text = f"[Документ {fname}]: {caption}"
            is_doc_edit = True
        except Exception as e: await msg.answer(f"❌ {str(e)[:150]}"); return
    elif msg.text:
        low = msg.text.lower()
        code_triggers = ["напиши код", "сделай игру", "напиши игру", "сделай сайт", "напиши сайт",
                        "создай игру", "сделай программу", "напиши программу", "создай сайт",
                        "игру на html", "игра на html", "код для сайта"]
        if any(w in low for w in code_triggers):
            await make_code(msg)
            return
        user_content = msg.text; save_text = msg.text
    else: return

    await save_message(uid, "user", save_text)
    history.append({"role": "user", "content": user_content})
    history = trim_history_by_chars(history)
    await bot.send_chat_action(msg.chat.id, "typing")
    now = datetime.now().strftime("%H:%M МСК, %d.%m.%Y")
    personal = PERSONAL_PROMPTS.get(uid, "")
    tone = await get_user_tone(uid)
    tone_add = f"\n\nТОН: {tone}" if tone else ""
    full_prompt = f"[Сейчас: {now}]\n\n" + SYSTEM_PROMPT + personal + tone_add
    try:
        r = await client.chat.completions.create(
            model="qwen/qwen3.8-27b",
            messages=[{"role": "system", "content": full_prompt}, *history],
            temperature=0.7, max_tokens=1200
        )
        answer = r.choices[0].message.content
        if not answer or not answer.strip():
            await msg.answer("...")
            return
        await save_message(uid, "assistant", answer)
        if is_doc_edit and doc_info:
            fb, fn, ext = doc_info
            await handle_document_edit(msg, answer, fb, fn, ext, uid)
        else:
            use_file = (mode == "file") or (len(answer) > 4000)
            if use_file:
                ext_out = detect_extension(answer)
                fname_out = f"light_answer.{ext_out}"
                code = extract_code(answer) if ext_out in CODE_EXTENSIONS else answer
                comment = await get_file_comment(f"файл .{ext_out}", "код", uid)
                await msg.answer_document(BufferedInputFile(code.encode("utf-8"), filename=fname_out), caption=comment)
            else:
                parts = split_message(answer)
                if len(parts) == 1: await msg.answer(parts[0])
                else:
                    for i, p in enumerate(parts, 1): await msg.answer(f"📄 Часть {i}/{len(parts)}\n\n{p}")
            if mode == "voice":
                await bot.send_chat_action(msg.chat.id, "record_voice")
                try:
                    vf = await text_to_voice(answer)
                    if vf: await msg.answer_voice(FSInputFile(vf))
                except Exception as e: logging.error(f"TTS: {e}")
    except Exception as e:
        logging.error(f"Ошибка: {e}")
        err = str(e)
        if "message text is empty" in err.lower():
            await msg.answer("...")
        else:
            await msg.answer(f"❌ {err[:300]}")

# === ВЕБ-СЕРВЕР ===
async def handle(request): return web.Response(text="Bot is running!")

async def main():
    global db_pool
    logging.basicConfig(level=logging.INFO)
    db_pool = await asyncpg.create_pool(DATABASE_URL, min_size=1, max_size=5)
    await init_db()
    logging.info("База данных подключена")
    if hf_client: logging.info("HF клиент ок")
    if pexafy_client: logging.info("Pexafy клиент ок")
    app = web.Application(); app.router.add_get("/", handle)
    runner = web.AppRunner(app); await runner.setup()
    port = int(os.environ.get("PORT", 8080))
    site = web.TCPSite(runner, "0.0.0.0", port); await site.start()
    logging.info(f"Web server port {port}")
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
